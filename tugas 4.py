"""Generator dashboard penjualan kafe (data simulasi, Excel + HTML).
Jalankan: pip install openpyxl; python generate_dashboard_kafe.py
"""
from pathlib import Path
from datetime import date, timedelta
import random, json, zlib, base64
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).resolve().parent
RNG = random.Random(2026)  # seed tetap agar data konsisten
PRODUCTS = [
    ("Kopi Susu", "Minuman",25000,10500), ("Americano","Minuman",22000,8000),
    ("Matcha Latte","Minuman",32000,14500), ("Teh Lemon","Minuman",18000,6500),
    ("Roti Bakar","Makanan",24000,10500), ("Croissant","Makanan",28000,12500),
    ("Nasi Ayam","Makanan",38000,19500), ("Brownies","Snack",26000,11000),
    ("Kentang Goreng","Snack",23000,9500),
]
BRANCHES = ["Makassar","Gowa","Maros","Parepare"]
CHANNELS = ["Dine-in","Takeaway","Delivery"]
MONTH_NAMES = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep"]

def generate_data(n=480):
    records=[]
    for i in range(1,n+1):
        product, category, price, cost = RNG.choice(PRODUCTS)
        month=RNG.randint(1,9)
        day=RNG.randint(1,28)
        qty=RNG.randint(1,5)
        discount=RNG.choice([0,0,0,5,10,15])
        sales=round(qty*price*(1-discount/100))
        profit=sales-qty*cost
        records.append(dict(id=f"KFE-{i:04d}",date=f"2026-{month:02d}-{day:02d}",year=2026,
            month=month,monthname=MONTH_NAMES[month-1],branch=RNG.choice(BRANCHES),
            channel=RNG.choice(CHANNELS),product=product,category=category,qty=qty,
            price=price,discount=discount,cost=cost,sales=sales,profit=profit,
            margin=round(100*profit/sales,2)))
    return records

def create_excel(records, path):
    wb=Workbook()
    ws=wb.active; ws.title="Data Transaksi"
    fields=["id","date","year","month","monthname","branch","channel","product",
            "category","qty","price","discount","cost","sales","profit","margin"]
    headers=["ID Transaksi","Tanggal","Tahun","Bulan","Nama Bulan","Cabang","Kanal",
             "Produk","Kategori","Jumlah","Harga Satuan","Diskon (%)","Biaya Satuan",
             "Penjualan (Rp)","Laba (Rp)","Margin (%)"]
    ws.append(headers)
    for r in records:
        ws.append([date.fromisoformat(r[f]) if f=="date" else r[f] for f in fields])
    ws.freeze_panes="A2"
    ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='12304A');cell.font=Font(color='FFFFFF',bold=True)
    for col,width in {'A':17,'B':15,'C':10,'D':10,'E':13,'F':17,'G':17,'H':20,
                      'I':15,'J':12,'K':17,'L':15,'M':18,'N':20,'O':18,'P':16}.items():
        ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):
        row[1].number_format='dd mmm yyyy'
        for idx in [10,12,13,14]:row[idx].number_format='"Rp" #,##0'
        row[15].number_format='0.00"%"'
    summary=wb.create_sheet('Ringkasan')
    summary.append(['DASHBOARD PENJUALAN KAFE | DATA SIMULASI 2026'])
    summary['A1'].font=Font(size=17,bold=True,color='12304A')
    summary.append(['Indikator','Nilai'])
    summary.append(['Total Penjualan',sum(r['sales'] for r in records)])
    summary.append(['Total Laba',sum(r['profit'] for r in records)])
    summary.append(['Total Transaksi',len(records)])
    total_sales=sum(r['sales'] for r in records)
    summary.append(['Margin (%)',100*sum(r['profit'] for r in records)/total_sales if total_sales else 0])
    summary['B3'].number_format=summary['B4'].number_format='"Rp" #,##0'
    summary['B6'].number_format='0.0"%"'
    summary.column_dimensions['A'].width=34;summary.column_dimensions['B'].width=23
    for c in summary[2]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='12304A')
    trend=wb.create_sheet('Analisis')
    trend.append(['Bulan','Penjualan','Laba'])
    for month in range(1,10):
        rows=[r for r in records if r['month']==month]
        trend.append([MONTH_NAMES[month-1],sum(r['sales'] for r in rows),sum(r['profit'] for r in rows)])
    trend.append([])
    trend.append(['Kategori','Penjualan','Laba'])
    for category in sorted(set(r['category'] for r in records)):
        rows=[r for r in records if r['category']==category]
        trend.append([category,sum(r['sales'] for r in rows),sum(r['profit'] for r in rows)])
    trend.append([])
    trend.append(['Cabang','Penjualan','Laba'])
    for branch in BRANCHES:
        rows=[r for r in records if r['branch']==branch]
        trend.append([branch,sum(r['sales'] for r in rows),sum(r['profit'] for r in rows)])
    trend.append([])
    trend.append(['Kanal','Penjualan','Laba'])
    for channel in CHANNELS:
        rows=[r for r in records if r['channel']==channel]
        trend.append([channel,sum(r['sales'] for r in rows),sum(r['profit'] for r in rows)])
    for col in 'ABC':trend.column_dimensions[col].width=23
    for row in trend.iter_rows():
        for cell in row[1:3]:cell.number_format='"Rp" #,##0'
    for start,end,title,anchor,kind in [(1,10,'Tren Bulanan','E2','line'),(12,15,'Penjualan per Kategori','E20','bar'),
                                         (17,21,'Penjualan per Cabang','N2','bar'),(23,26,'Penjualan per Kanal','N20','bar')]:
        chart=LineChart() if kind=='line' else BarChart()
        chart.title=title;chart.y_axis.title='Rupiah';chart.x_axis.title='Kelompok'
        chart.add_data(Reference(trend,min_col=2,max_col=3,min_row=start,max_row=end),titles_from_data=True)
        chart.set_categories(Reference(trend,min_col=1,min_row=start+1,max_row=end))
        chart.width=18;chart.height=9
        trend.add_chart(chart,anchor)
    wb.save(path)

# Template HTML yang sama dengan dashboard interaktif sebelumnya, disimpan terkompresi.
TEMPLATE_PACKED = """eNqdWOty27gVfhVG2QzJmqIudhKbEpnNOttpduN2x/E/jyYBSUhCBF4GAGWpWs70HTqd6Zv0fx9ln6TnAKRujrLeesYScIhzwbl851DjZ2mRqHVJrbnKeDTGT4uTfBZ2WNqBPSVpNM6oIlYyJ0JSFXYqNe1edhpqTjIadpaMPpSFUB0rKXJFczj1wFI1D1O6ZAnt6o3HcqYY4V2ZEE7DAYhQTHEavSNyHhdEpNbPZEqtYX/4atwzj8YyEaxUlhRJ2JkrVcqg10vS3P8iU8rZUvg5Vb28zHpongIySO0ZJmBWa5ARF+l6kxExY3nQH8UkWcxEUeVp8Hw6nL6aktEUbA4GF+XKeivAvlFS8EIEzwcXw+RiWKMLqNjs8w2G5/0L0px7mDNFRyVJU5bPgiGKefming+2Kq2+dVmu6oywHGgr4wzQ1++Xq1FziFSqOJBRS5ooVuT7eo2muBBgT1eQlFUyGAxBSMs4uNxK7MaFUkVmSHGx6so5SYsHsAYYLOSymmsMhrU/ZVxRITcpkyUn62DK6Wo0I6URj7vug4AtfoBlHGzz4goU5Jut7r5WhLYFA5AuC85S63lC0td0cGT0azx6dK/aTyADdibMBEtH+NFVNAOKol3wd5XlMhC0pEQ5F95gKlxjJnpMCziIE72CCF8duufIff2W0ZJKFPlsqz/mRbLQqdGV7O80GJ7vfKuKMrjSjJh0T7IZTLXg31h7ueXdzCmbzVVwjslQJyRfEqmTpKEPXyFdkZjTTWM5SOSklDRoF6Mmofr9F7VKPTXfxuRqd982HXaRoeeUTPsjRVeqSzib5QGnU1UD+2MX1t9nNGXE2aXv5UswzN2YoHmtI7559/qxlJd/TEpdQ23rmh73DDJhbRuUogK+B9Fv//6XtcOTX2j+pSKAZxpZgGkQjVO2jH4ieUUE++0f//xIS9AUU6Fhx/rvf4AbYE2yrOJEMismokLqR5pVgHaME2alIDGzbquSkfm4h/KMOWgCVjkAjyleKwEZMuw05QXQxElMeXRNYoDYcSzwJBaTxdKwEwuSJ3ONX5oIC3O84foZvDErBDvmSwx9/S3OnPBHbHOS55Sf5voBPJAfc2VQEF+10cCBVeQJZ8ki7AgKzcJxO9EtLizjgnHPHNP82kWPfKVToaPDtE/qRHeFInwX0bGpV22UhH5igF/ToiYoJ0R8AO/vc5eimDL1FHYIkCQLyfa5dXk9RfmNxo59VoMmj1l3vtkXomujs/UXZPs5WERzS8cJHNIDwiMOYDCoojUqOI9tvWdojxXuC98VTwnlsc2+J6hpMpKZoPxxXU19PEGTqZn/V4+pi6dcyNTKN9WY/b6yW8DgBZGgbJs51h0VbSm0alFBXija2U8dDXOQXUs4z6Fv44CAAxN2Avgy6KcELqP372BcmuvlHfhthndq9q0nm+0vokirxXbb3L99uC2tlqIrRW96qKrXqkXU1WanMAIyDSHKIHGvMfCRcxpcNIMZDIlSAYwqEn769O7t3dtPP338218/fRqZB6IqwzyM7NvSss9uiJr7uhk5ueur4kOB0+NHJcC5jg2N4v072x1xgBhTIeGmbsRMGeWpDO9tkyS2Z7dIiUsTUVhpRLMn0OqFYxgXVjFtuN2NIVEewqBcZTDa+jOqfuQUlz+s36fOwmhfEg66fB8G0gfoFsrB6/kZKR0RRuJ+MXHdiS9hRnYc4sVuGOHUDXpIGIZ2XmELst+Qbhw0VyOuz/VVr4usJII6DT12XXdEuc/AfPGXu5sPoT0uSo2gYEIFOdOJdLMa9ww5ss/QNm3KMow+H53+brOsO9F3mwXaYXzxBq8AbZA62BO9ZXfgDU653ttonsDmAGB27QbLeqv4s+t/KVju2La2GFoDOH1Gw6oE39BRPa1yA/1Nr9gYn/sQiR9JMncWYXTa6b6xH2UbeY67E0hmM0FnSBTFg/QWdO1uMEhZiFe7AU+4e/EWGG88aA4tQwgXXU9GbOo8y/w5kc7SdTMfbVx6G91sgr5n2kbQr038V2GGRsLR0crXZ85CYRZAMIeRYla1oKoSuYUJk/lwKwRL5yhDwIi9oJD7/qQbw0cQ3w8mRnKXbJd7t08FeXBY6rV39zDV2kwuCRMy/IqDPN3IZagP6GxZHZtwmBcrNOgpqSHncCnMDeRw0a2mVO9ZOnF3Sz+l2AfXEJsdUQfsGrfOqVxgqett8IoellywMffQGwiZDO4NJbC3AGebkwc3Xe086e0G4Gv9imc/H5LXCbm0PTNNH1Frr9WAgHlSuIn816QPrsjrS3osvaHWk9ozBSWDDVRKCQu2pIESFfUQVgGC87eyBKi9JXAsmEK1U6/kFQwXwKKKgitWBhuIEEfdsnFRkIRR4jdu8jXpzA4AcgF/ncQv8XU/9dduXdeefmUHxjXIY1pEKy0ATHGWPXj7wD/MhT+zFU2dgXtmW18g7Pi3l5xtsTbpiOkXaqg0bRHRsoEBCr1vjSDw7PdQ4Ndfm8zTOAsJ+3sMAKFGvQ54iEb4gqZVQh1HegJqT541xev13abUTx0zT+Hc6JRaW0uywTnwsnXd/EKCXjb5dprPiP4KY5NLpznNWHrMiRfgNJ9BSz3VSU9KNNPqkURzgzcQ/T8Zk3oNPB4mwgt7dH9v6+lz23E9m7Oc2hPv3t7Ni4dNGl6+9PN2ykNS28vbZ+1odtDT9cPJtpFABWpMRPwxqIgFCR9D6Mqnb4xD2WkPQnqr7VgnqazgzdDksG/d0JxkJeMLmP3O+3vnYOpUJCO+fVqrmalA767Da60S3quo0/fOwbfNWPHZTIAptG/hsxQ6r0q3eyyzQ4rx3SGtxIkwUYfExpEHREg60XaaR/QmHZsHOCzu9f66LfnR7pe5XjMu6l8d/wcCAEPs"""

def create_html(records,path):
    template=zlib.decompress(base64.b64decode(TEMPLATE_PACKED)).decode('utf-8')
    content=template.replace('__DATA_JSON__',json.dumps(records,ensure_ascii=False))
    path.write_text(content,encoding='utf-8')

if __name__=='__main__':
    rows=generate_data()
    xlsx=OUT/'Dashboard_Penjualan_Kafe_Baru.xlsx'
    html=OUT/'Dashboard_Penjualan_Kafe_Interaktif.html'
    create_excel(rows,xlsx);create_html(rows,html)
    print('Berhasil:',xlsx,html)
    print('Transaksi:',len(rows),'Total penjualan:',sum(r['sales'] for r in rows))
